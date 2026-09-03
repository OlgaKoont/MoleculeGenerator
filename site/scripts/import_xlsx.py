from __future__ import annotations

import re
from pathlib import Path
import sys

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from mappings import (
    ALIASES,
    ARCHITECTURE_IDS,
    COMPATIBILITY_GROUP,
    GLOSSARY_FOR_TOOL,
    INCOMPATIBLE_GROUPS,
    OUTPUT_MODALITY,
    PAPER_TITLE_CONFIDENCE,
    PAPER_TITLES,
    RELATED,
    TASK_IDS,
    TOOL_SLUGS,
    WORKFLOW_IDS,
)


ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "content"
XLSX_CANDIDATES = [
    Path(__file__).resolve().parents[2]
    / "ML_generators_protein_antibody_landscape_RU (1).xlsx",
]


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text == "" or text.upper() == "NR":
        return None
    if text.lower().startswith("nr "):
        return None
    if text in {"NR — no verified official maintained repository found", "NR — official maintained repository not verified", "NR — no verified turnkey official software package"}:
        return None
    return text


def looks_like_url(value: str | None) -> str | None:
    value = clean(value)
    if not value:
        return None
    if value.startswith("http://") or value.startswith("https://"):
        return value
    return None


def looks_like_doi(value: str | None) -> str | None:
    value = clean(value)
    if not value:
        return None
    if value.lower().startswith("10."):
        return value
    return None


def parse_flag(value) -> bool:
    return str(value).strip() == "1"


def split_limitations(text: str | None) -> list[str]:
    if not text:
        return []
    parts = re.split(r";\s*", text)
    return [p.strip().rstrip(".") + "." for p in parts if p.strip()]


def generator_class_tokens(raw: str | None) -> list[str]:
    if not raw:
        return []
    mapping = [
        ("1", "sequence-generator"),
        ("2", "inverse-folding"),
        ("3", "backbone-structure-generator"),
        ("4", "joint-sequence-structure"),
        ("5", "target-conditioned-binder"),
        ("6", "mutation-optimizer"),
        ("7", "humanization"),
    ]
    found = []
    for key, token in mapping:
        if key in raw:
            found.append(token)
    return found or [raw]


def classify_repo(url: str | None, availability: str | None, name: str) -> tuple[str | None, str | None, str | None]:
    repo = None
    docs = None
    service = None
    if url:
        if "github.com" in url:
            repo = url
        elif "genscript.com" in url or "service" in url.lower() or "tool.html" in url:
            service = url
        else:
            docs = url
    return repo, docs, service


def incompatible_for(slug: str) -> list[str]:
    group = COMPATIBILITY_GROUP.get(slug)
    if not group:
        return []
    other_groups = INCOMPATIBLE_GROUPS.get(group, [])
    return sorted(
        other
        for other, g in COMPATIBILITY_GROUP.items()
        if g in other_groups and other != slug
    )


def workflow_breakdown(row: list, slug: str) -> list[dict]:
    primary = clean(row[17])
    scorer = clean(row[18])
    predictor = clean(row[19])
    integrated = clean(row[20])
    external = clean(row[21])
    items = []
    if primary:
        items.append(
            {
                "name": primary,
                "role": "primary_generator",
                "architecture_class": clean(row[4]),
                "integrated_or_external": "integrated",
                "required_or_optional": "required",
                "influences_generation_or_post_hoc": "generation",
                "runtime_included_or_excluded": "included",
            }
        )
    if scorer:
        items.append(
            {
                "name": scorer,
                "role": "integrated_scorer",
                "architecture_class": None,
                "integrated_or_external": "integrated",
                "required_or_optional": "optional",
                "influences_generation_or_post_hoc": "post_hoc",
                "runtime_included_or_excluded": "unknown",
            }
        )
    if predictor and predictor.lower() not in {"no", "no dedicated structure/complex predictor"}:
        role = "integrated_complex_predictor" if "complex" in predictor.lower() else "integrated_structure_predictor"
        if predictor.lower().startswith("no"):
            pass
        else:
            items.append(
                {
                    "name": predictor,
                    "role": role,
                    "architecture_class": None,
                    "integrated_or_external": "integrated",
                    "required_or_optional": "optional",
                    "influences_generation_or_post_hoc": "post_hoc",
                    "runtime_included_or_excluded": "included",
                    "notes": "Predictor/ranker, not a generator.",
                }
            )
    if external:
        items.append(
            {
                "name": external,
                "role": "mandatory_external_dependency",
                "architecture_class": None,
                "integrated_or_external": "external",
                "required_or_optional": "required",
                "influences_generation_or_post_hoc": "post_hoc",
                "runtime_included_or_excluded": "excluded",
            }
        )
    items.append(
        {
            "name": "Experimental binding/function assays",
            "role": "experimental_validation",
            "architecture_class": None,
            "integrated_or_external": "external",
            "required_or_optional": "required",
            "influences_generation_or_post_hoc": "post_hoc",
            "runtime_included_or_excluded": "excluded",
            "notes": "Only assays establish measured affinity or function.",
        }
    )
    return items


def import_tools() -> tuple[list[dict], list[dict]]:
    xlsx = next(p for p in XLSX_CANDIDATES if p.exists())
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Генераторы"]
    tools = []
    papers = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        vals = list(row) + [None] * (36 - len(row))
        name = clean(vals[0])
        if not name:
            continue
        slug = TOOL_SLUGS[name]
        paper_id = f"paper-{slug}"
        repo_url = looks_like_url(vals[28])
        repo, docs, service = classify_repo(repo_url, clean(vals[30]), name)
        doi = looks_like_doi(vals[27])
        paper_url = looks_like_url(vals[26])
        pub_status = clean(vals[25])
        preprint = parse_flag(vals[35]) or (pub_status or "").lower().find("preprint") >= 0
        evidence_class = clean(vals[22])
        prospective = parse_flag(vals[33])
        evidence = {
            "evidence_level": evidence_class,
            "retrospective_validation": None if prospective and evidence_class and evidence_class.startswith("3") else clean(vals[23]),
            "prospective_binding_validation": clean(vals[24]) if prospective else None,
            "prospective_functional_validation": clean(vals[24]) if prospective and evidence_class and ("4" in evidence_class or "5" in evidence_class) else None,
            "developability_validation": clean(vals[24]) if evidence_class and "5" in evidence_class else None,
            "tested_design_count": clean(vals[23]),
            "success_criterion": clean(vals[24]),
            "independent_validation": None,
            "evidence_summary": clean(vals[24]),
            "prospective_flag": prospective,
            "direct_antibody_flag": parse_flag(vals[34]),
            "preprint_flag": preprint,
        }
        rel = RELATED.get(slug, {})
        review = []
        title = PAPER_TITLES.get(slug)
        if slug in PAPER_TITLE_CONFIDENCE:
            review.append("Bibliographic title must be verified against the PDF.")
        if not doi and not paper_url:
            review.append("Primary paper link missing.")
        if not repo and not service:
            review.append("Official repository/service not verified.")
        if PAPER_TITLE_CONFIDENCE.get(slug) == "review-queue" or not title:
            title_kind = "missing" if not title else "interpretation"
        else:
            title_kind = "author-reported"

        tool = {
            "identity": {
                "id": slug,
                "name": name,
                "slug": slug,
                "aliases": ALIASES.get(slug, []),
                "version": None,
                "release_year": clean(vals[1]),
                "last_verified": "2026-08-13",
                "official_repository": repo,
                "official_documentation": docs,
                "official_service": service,
                "license": clean(vals[29]),
                "commercial_use_status": clean(vals[30]),
                "publication_status": pub_status,
            },
            "classification": {
                "task_ids": TASK_IDS.get(slug, []),
                "workflow_ids": WORKFLOW_IDS.get(slug, []),
                "architecture_ids": ARCHITECTURE_IDS.get(slug, []),
                "generator_classes": generator_class_tokens(clean(vals[3])),
                "output_modalities": OUTPUT_MODALITY.get(slug, []),
                "supported_formats": clean(vals[8]),
                "de_novo_or_optimization": clean(vals[9]),
                "general_protein_relevance": clean(vals[13]),
                "antibody_relevance": clean(vals[14]),
                "VHH_relevance": clean(vals[15]),
                "mutation_relevance": clean(vals[16]),
                "compatibility_group": COMPATIBILITY_GROUP.get(slug),
            },
            "generation": {
                "exact_generated_object": clean(vals[5]) or "NR",
                "required_input": clean(vals[6]),
                "conditioning": clean(vals[7]),
                "fixed_components": None,
                "editable_components": None,
                "generated_components": clean(vals[5]),
                "support_for_single_mutations": clean(vals[10]),
                "support_for_combinatorial_mutations": clean(vals[11]),
                "support_for_multi_objective_optimization": clean(vals[12]),
                "candidate_generation_mechanism": clean(vals[4]),
            },
            "architecture": {
                "base_model": None,
                "model_family": clean(vals[4]),
                "model_type": clean(vals[3]),
                "input_representation": None,
                "node_features": None,
                "edge_or_pair_features": None,
                "encoder": None,
                "decoder": None,
                "latent_or_generative_state": None,
                "equivariance": None,
                "conditioning_mechanism": clean(vals[7]),
                "generation_mechanism": clean(vals[4]),
                "sampling_procedure": None,
                "training_objectives": None,
                "losses": None,
                "architectural_modifications": None,
                "computational_complexity_drivers": None,
                "output": clean(vals[5]),
                "internal_scorer": clean(vals[18]),
                "structure_generation": clean(vals[19]),
                "summary_notes": [],
                "mathematical_details": None,
                "source_label": "Author-reported (matrix)",
            },
            "workflow_components": {
                "primary_generator": clean(vals[17]),
                "integrated_generators": [],
                "integrated_scorers": [clean(vals[18])] if clean(vals[18]) else [],
                "integrated_structure_predictor": clean(vals[19]),
                "integrated_complex_predictor": None,
                "integrated_relaxation": None,
                "mandatory_external_dependencies": [clean(vals[21])] if clean(vals[21]) else [],
                "optional_external_dependencies": [clean(vals[20])] if clean(vals[20]) else [],
                "breakdown": workflow_breakdown(vals, slug),
            },
            "training": {
                "training_data": None,
                "preprocessing": None,
                "negative_examples": None,
                "fine_tuning": None,
                "label_requirements": None,
                "reported_training_hardware": None,
                "reported_training_time": None,
            },
            "inference": {
                "candidate_count": None,
                "batch_size": None,
                "generation_steps": None,
                "recycles": None,
                "sequence_or_complex_size": None,
                "reported_runtime": None,
                "runtime_boundary": None,
                "CPU": None,
                "GPU": None,
                "VRAM": None,
                "RAM": None,
                "disk": None,
                "database_requirements": None,
            },
            "evidence": evidence,
            "limitations": {
                "scientific_limitations": split_limitations(clean(vals[32])),
                "architectural_limitations": [],
                "dataset_biases": [],
                "likely_failure_modes": [],
                "reproducibility_limitations": [],
                "deployment_limitations": [clean(vals[31])] if clean(vals[31]) else [],
            },
            "relations": {
                "related_tools": rel.get("related_tools", []),
                "alternatives": rel.get("alternatives", []),
                "complementary_tools": rel.get("complementary_tools", []),
                "incompatible_comparisons": incompatible_for(slug),
                "paper_ids": [paper_id],
                "source_ids": ["landscape-xlsx-2026-08-13"],
                "dependency_ids": [],
                "glossary_ids": GLOSSARY_FOR_TOOL.get(slug, []),
            },
            "detailed_architecture": {},
            "review_queue": review,
        }
        paper = {
            "id": paper_id,
            "slug": slug,
            "title": title,
            "authors": None,
            "year": clean(vals[1]),
            "venue": pub_status,
            "publication_status": pub_status,
            "doi": doi,
            "url": paper_url,
            "preprint_url": paper_url if preprint else None,
            "tool_ids": [slug],
            "main_task": clean(vals[2]),
            "claimed_novelty": None,
            "model_architecture": clean(vals[4]),
            "input_representation": None,
            "encoder_decoder": None,
            "generative_state": clean(vals[5]),
            "conditioning": clean(vals[7]),
            "sampling": None,
            "training_objectives": None,
            "training_dataset": None,
            "preprocessing": None,
            "inference": None,
            "baselines": None,
            "retrospective_benchmarks": None if prospective else clean(vals[24]),
            "prospective_experimental_validation": clean(vals[24]) if prospective else None,
            "tested_designs": clean(vals[23]),
            "success_criterion": clean(vals[24]),
            "main_results": clean(vals[24]),
            "limitations": clean(vals[32]),
            "demonstrates": [clean(vals[24])] if clean(vals[24]) else [],
            "does_not_demonstrate": [
                "Predicted confidence, likelihood, pLDDT, docking or developability scores are not experimental affinity."
            ],
            "claims": [
                {
                    "text": title or "Primary bibliographic title",
                    "kind": title_kind,
                    "locator": "DOI/URL from landscape matrix",
                },
                {
                    "text": f"Publication status: {pub_status or 'NR'}",
                    "kind": "author-reported",
                    "locator": "landscape matrix, publication status column",
                },
            ],
            "review_queue": review,
        }
        tools.append(tool)
        papers.append(paper)
    return tools, papers


def dump_yaml(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=88),
        encoding="utf-8",
    )


def main() -> None:
    tools, papers = import_tools()
    for tool in tools:
        dump_yaml(CONTENT / "tools" / f"{tool['identity']['slug']}.yaml", tool)
    for paper in papers:
        dump_yaml(CONTENT / "papers" / f"{paper['slug']}.yaml", paper)
    dump_yaml(
        CONTENT / "meta" / "source.yaml",
        {
            "id": "landscape-xlsx-2026-08-13",
            "name": "ML_generators_protein_antibody_landscape_RU",
            "verified": "2026-08-13",
            "path": "ML_generators_protein_antibody_landscape_RU (1).xlsx",
            "note": "Primary curated matrix for generators, evidence flags, DOI and repository URLs.",
        },
    )
    print(f"wrote {len(tools)} tools and {len(papers)} papers")


if __name__ == "__main__":
    main()
